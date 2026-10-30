import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

IMG_DIR = r'C:\xampp\htdocs\Attendai\extracted_doc_images'

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="none"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def build_complete_document():
    doc = docx.Document()

    # Page Setup: A4, 1-inch margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)

    # Base Normal Style: 12pt Calibri
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)

    def add_h14(text, space_before=14, space_after=4):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = 'Calibri'
        r.font.size = Pt(14)
        r.font.bold = True
        r.font.color.rgb = RGBColor(15, 23, 42)
        return p

    def add_h13(text, space_before=10, space_after=3):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = 'Calibri'
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = RGBColor(30, 41, 59)
        return p

    def add_body(text, bold_prefix=None, space_after=5):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = 'Calibri'
            r_pre.font.size = Pt(12)
            r_pre.font.bold = True
        r = p.add_run(text)
        r.font.name = 'Calibri'
        r.font.size = Pt(12)
        return p

    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = 'Calibri'
            r_pre.font.size = Pt(12)
            r_pre.font.bold = True
        r = p.add_run(text)
        r.font.name = 'Calibri'
        r.font.size = Pt(12)
        return p

    def create_table(headers, rows, col_widths=None):
        table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        set_table_borders(table)

        # Header Row
        hdr_cells = table.rows[0].cells
        for i, header_text in enumerate(headers):
            hdr_cells[i].text = header_text
            set_cell_background(hdr_cells[i], "0F172A") # Slate 900
            set_cell_margins(hdr_cells[i], top=110, bottom=110, left=130, right=130)
            p = hdr_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.name = 'Calibri'
                run.font.size = Pt(11)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)

        # Body Rows
        for r_idx, row_data in enumerate(rows):
            row_cells = table.rows[r_idx + 1].cells
            bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
            for c_idx, cell_value in enumerate(row_data):
                row_cells[c_idx].text = str(cell_value)
                set_cell_background(row_cells[c_idx], bg_color)
                set_cell_margins(row_cells[c_idx], top=90, bottom=90, left=130, right=130)
                p = row_cells[c_idx].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    run.font.name = 'Calibri'
                    run.font.size = Pt(11)
                    run.font.color.rgb = RGBColor(51, 65, 85)

        if col_widths:
            for row in table.rows:
                for idx, w in enumerate(col_widths):
                    row.cells[idx].width = Inches(w)

        p_sp = doc.add_paragraph()
        p_sp.paragraph_format.space_before = Pt(2)
        p_sp.paragraph_format.space_after = Pt(4)
        return table

    def add_embedded_image(img_filename, caption, width=Inches(5.8), height=None):
        img_path = os.path.join(IMG_DIR, img_filename)
        if os.path.exists(img_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(8)
            p_img.paragraph_format.space_after = Pt(3)
            p_img.paragraph_format.keep_with_next = True
            run = p_img.add_run()
            if height:
                run.add_picture(img_path, width=width, height=height)
            else:
                run.add_picture(img_path, width=width)

            # Caption
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(2)
            p_cap.paragraph_format.space_after = Pt(12)
            r_cap = p_cap.add_run(caption)
            r_cap.font.name = 'Calibri'
            r_cap.font.size = Pt(11)
            r_cap.font.italic = True
            r_cap.font.color.rgb = RGBColor(71, 85, 105)

    def add_screenshot_figure(img_filename, fig_num, title, url):
        p_u = doc.add_paragraph()
        p_u.paragraph_format.space_before = Pt(6)
        p_u.paragraph_format.space_after = Pt(1)
        p_u.paragraph_format.keep_with_next = True
        r_t = p_u.add_run(f"{title}\n")
        r_t.font.name = 'Calibri'
        r_t.font.size = Pt(12)
        r_t.font.bold = True
        r_t.font.color.rgb = RGBColor(15, 23, 42)
        r_u = p_u.add_run(f"URL: {url}")
        r_u.font.name = 'Calibri'
        r_u.font.size = Pt(11)
        r_u.font.italic = True
        r_u.font.color.rgb = RGBColor(37, 99, 235)

        img_path = os.path.join(IMG_DIR, img_filename)
        if os.path.exists(img_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(4)
            p_img.paragraph_format.space_after = Pt(2)
            p_img.paragraph_format.keep_with_next = True
            p_img.add_run().add_picture(img_path, width=Inches(5.8))

            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(2)
            p_cap.paragraph_format.space_after = Pt(14)
            r_cap = p_cap.add_run(f"Figure {fig_num}: {title}")
            r_cap.font.name = 'Calibri'
            r_cap.font.size = Pt(11)
            r_cap.font.italic = True
            r_cap.font.color.rgb = RGBColor(71, 85, 105)

    # ==================== PAGE 1: COVER PAGE ====================
    p_t = doc.add_paragraph()
    p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t.paragraph_format.space_before = Pt(24)
    p_t.paragraph_format.space_after = Pt(12)
    r_t = p_t.add_run("AttendAI: AI-Powered Smart Attendance &\nClassroom Management System")
    r_t.font.name = 'Calibri'
    r_t.font.size = Pt(18)
    r_t.font.bold = True
    r_t.font.color.rgb = RGBColor(15, 23, 42)

    p_sb = doc.add_paragraph()
    p_sb.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sb.paragraph_format.space_after = Pt(16)
    r_sb = p_sb.add_run("Project Report\nSubmitted for the subject of\n221601907 Capstone Project – II\nIntegrated Master of Science (IT) (IMSc (IT))\nSemester IX")
    r_sb.font.name = 'Calibri'
    r_sb.font.size = Pt(12)
    r_sb.font.color.rgb = RGBColor(51, 65, 85)

    p_sbm = doc.add_paragraph()
    p_sbm.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sbm.paragraph_format.space_before = Pt(16)
    p_sbm.paragraph_format.space_after = Pt(4)
    r_s1 = p_sbm.add_run("Submitted By\n")
    r_s1.font.bold = True
    r_s1.font.size = Pt(12)
    r_s2 = p_sbm.add_run("Patel Abhi Kiranbhai (2026001)\nPatel Mahir Tarpankumar (001)")
    r_s2.font.size = Pt(12)

    p_gd = doc.add_paragraph()
    p_gd.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_gd.paragraph_format.space_before = Pt(16)
    p_gd.paragraph_format.space_after = Pt(4)
    r_g1 = p_gd.add_run("Under the Guidance of\n")
    r_g1.font.bold = True
    r_g1.font.size = Pt(12)
    r_g2 = p_gd.add_run("Prof. Sadhana Mishra\nFCAIT, GLS University")
    r_g2.font.size = Pt(12)

    # University Logos
    p_lg = doc.add_paragraph()
    p_lg.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_lg.paragraph_format.space_before = Pt(20)
    p_lg.paragraph_format.space_after = Pt(8)
    logo1 = os.path.join(IMG_DIR, 'page_1_img_1.jpeg')
    logo2 = os.path.join(IMG_DIR, 'page_1_img_2.jpeg')
    if os.path.exists(logo1):
        p_lg.add_run().add_picture(logo1, width=Inches(2.2))
        p_lg.add_run("   ")
    if os.path.exists(logo2):
        p_lg.add_run().add_picture(logo2, width=Inches(1.8))

    p_un = doc.add_paragraph()
    p_un.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_un.paragraph_format.space_before = Pt(12)
    p_un.paragraph_format.space_after = Pt(2)
    r_un = p_un.add_run("Faculty of Computer Applications & IT (FCAIT)\nIMSc (IT) Programme\nAhmedabad\nAcademic Year 2026–2027")
    r_un.font.name = 'Calibri'
    r_un.font.size = Pt(12)
    r_un.font.bold = True
    r_un.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_page_break()

    # ==================== PAGE 2: CERTIFICATE ====================
    p_c1 = doc.add_paragraph()
    p_c1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_c1.paragraph_format.space_before = Pt(16)
    r_c1 = p_c1.add_run("GLS UNIVERSITY\nFaculty of Computer Applications & IT\nIMSc (IT) Programme Ahmedabad\n\n")
    r_c1.font.size = Pt(14)
    r_c1.font.bold = True

    r_c2 = p_c1.add_run("CERTIFICATE\n")
    r_c2.font.size = Pt(16)
    r_c2.font.bold = True
    r_c2.font.underline = True

    p_cert = doc.add_paragraph()
    p_cert.paragraph_format.space_before = Pt(24)
    p_cert.paragraph_format.line_spacing = 1.25
    r_cert = p_cert.add_run(
        "This is to certify that the following students of Semester IX Integrated Master of Science (IT) (IMSc (IT)) Programme, "
        "Faculty of Computer Applications & IT (FCAIT), GLS University, Ahmedabad, have successfully completed the subject project titled "
        "\"AttendAI: Smart Attendance & Classroom Management System\" for the course 221601907 Capstone Project – II - System Design and "
        "Implementation during the Academic Year 2026–2027."
    )
    r_cert.font.size = Pt(12)

    p_names = doc.add_paragraph()
    p_names.paragraph_format.space_before = Pt(20)
    r_nm = p_names.add_run("Students:\n1. Patel Abhi Kiranbhai (Enrollment ID: 2026001)\n2. Patel Mahir Tarpankumar (Enrollment ID: 001)")
    r_nm.font.size = Pt(12)
    r_nm.font.bold = True

    p_dt = doc.add_paragraph()
    p_dt.paragraph_format.space_before = Pt(50)
    p_dt.add_run("Date of Submission: ____________________\n\n\n\n")

    p_sg = doc.add_paragraph()
    p_sg.paragraph_format.space_before = Pt(40)
    r_sga = p_sg.add_run("Prof. Sadhana Mishra\t\t\t\t\t\tDr. Tripti Dodiya\nProject Guide\t\t\t\t\t\t\tDean, FCAIT")
    r_sga.font.size = Pt(12)
    r_sga.font.bold = True

    doc.add_page_break()

    # ==================== PAGE 3: INDEX (TABULAR FORMAT) ====================
    add_h14("3. Index / Table of Contents", space_before=10)

    index_table_data = [
        ["1", "Cover Page", "1"],
        ["2", "College Certificate", "2"],
        ["3", "Index / Table of Contents", "3"],
        ["4", "Brief of System", "4"],
        ["", "4.1 Technology Stack", "4"],
        ["", "4.2 Roles and Capabilities (Admin, Faculty, Student)", "4"],
        ["", "4.3 Attendance Verification Pipeline & Adaptive Biometrics", "5"],
        ["", "4.4 Security & Data Privacy", "5"],
        ["", "4.5 REST API Endpoints Specification", "6"],
        ["", "4.6 Project Objectives", "7"],
        ["", "4.7 Functional and Non-Functional Requirements", "7"],
        ["", "4.8 Scope & Limitations", "8"],
        ["", "4.9 Future Scope & Enhancements", "8"],
        ["", "4.10 Conclusion", "8"],
        ["5", "UML Diagrams", "9"],
        ["", "5.1 Use Case Diagram (Figure 1)", "9"],
        ["", "5.2 Class Diagram & Relational Foreign Keys (Figure 2)", "10"],
        ["", "5.3 Swimlane Activity Diagram: Attendance Marking (Figure 3)", "11"],
        ["", "5.4 Sequence Diagram: Session Start & Verification (Figure 4)", "12"],
        ["", "5.5 State Chart Diagrams: Session & Biometrics (Figures 5 & 6)", "13"],
        ["6", "Data Dictionary (Database Models & Schemas)", "14"],
        ["", "6.1 User Schema", "14"],
        ["", "6.2 Department Schema", "15"],
        ["", "6.3 Course / Branch Schema (NEW)", "16"],
        ["", "6.4 Subject Schema", "17"],
        ["", "6.5 Class Schema", "18"],
        ["", "6.6 Attendance Schema & Embedded Records", "19"],
        ["", "6.7 Relational Model & Referential Integrity (Foreign Keys)", "20"],
        ["7", "All System Screenshots with URLs", "21"],
        ["", "7.1 Admin Portal Interfaces", "22"],
        ["", "7.2 Faculty Portal Interfaces", "25"],
        ["", "7.3 Student Portal Interfaces", "28"],
        ["", "7.4 Reports, Analytics & PDF Export", "30"],
        ["8", "References", "32"]
    ]
    create_table(["Sr No.", "Topic / Section Title", "Page No."], index_table_data, [0.8, 4.6, 0.8])

    add_h14("List of Figures & Diagrams", space_before=16)
    figures_table_data = [
        ["Figure 1", "Use Case Diagram - Multi-Role Actor Interactions", "9"],
        ["Figure 2", "Class Diagram - Mongoose Schemas & Foreign Key Relationships", "10"],
        ["Figure 3", "Swimlane Activity Diagram - Attendance Marking Workflow", "11"],
        ["Figure 4", "Sequence Diagram - Session Start, Biometrics & Check-in", "12"],
        ["Figure 5", "State Chart Diagram - Attendance Session Lifecycle", "13"],
        ["Figure 6", "State Chart Diagram - Student Biometric Verification", "13"],
        ["Figure 7", "Admin Dashboard - Statistics & Operational Counters", "22"],
        ["Figure 8", "System Login Portal - Multi-Role Authentication", "22"],
        ["Figure 9", "User Management - Accounts, Roles & Access Control", "23"],
        ["Figure 10", "Department Management - Academic Departments (FCAIT)", "23"],
        ["Figure 11", "Subject Library - Master Curriculum Catalog", "24"],
        ["Figure 12", "All Classes - Active Divisions & Timetable Bindings", "24"],
        ["Figure 13", "Faculty Dashboard - Today's Schedule & Quick Actions", "25"],
        ["Figure 14", "My Classes - Assigned Divisions & Enrolled Rosters", "25"],
        ["Figure 15", "Master Timetable Schedule - Weekly Classroom Grid", "26"],
        ["Figure 16", "My Students - Enrolled Student Profiles & Face Status", "26"],
        ["Figure 17", "Live Attendance QR Hub - QR, Passcode & Radius Settings", "27"],
        ["Figure 18", "Subjects - Faculty Curriculum View", "27"],
        ["Figure 19", "Student Dashboard - Overall Attendance & Today's Schedule", "28"],
        ["Figure 20", "Student Registration - Enrollment ID & Biometric Enrollment", "28"],
        ["Figure 21", "Biometric QR Scanner - Multi-Pass Face Check & GPS", "29"],
        ["Figure 22", "My History - Student Attendance Verification Log", "29"],
        ["Figure 23", "Session Archives - Historical Attendance Logs with CSV Export", "30"],
        ["Figure 24", "Timetable Print Layout - Clean PDF Schedule Export", "31"]
    ]
    create_table(["Figure No.", "Figure Title & Description", "Page No."], figures_table_data, [1.1, 4.3, 0.8])

    doc.add_page_break()

    # ==================== PAGE 4: BRIEF OF SYSTEM ====================
    add_h14("4. Brief of System")
    add_body(
        "AttendAI is an enterprise-grade, cloud-deployed smart attendance and classroom management platform engineered to eliminate proxy attendance and manual administrative overhead in higher education institutions. The system synergistically combines multi-factor biometric and spatial verification: an encrypted rotating session token embedded in dynamic QR codes, a 4-digit session security pin, high-precision GPS geofencing, and client-side facial landmark verification using deep neural networks."
    )
    add_body(
        "The application is built on modern MERN stack architecture: a reactive single-page client built with React 19, Vite, and Tailwind CSS, interfacing with a Node.js and Express 5 backend deployed on Vercel, paired with a scalable cloud-hosted MongoDB Atlas replica set managed via Mongoose 9 schemas."
    )

    add_h14("4.1 Technology Stack")
    add_bullet("Frontend Framework: React 19, React Router 7, Vite 8, Tailwind CSS 4.")
    add_bullet("Biometrics & Machine Learning: face-api.js (SSD MobileNet V1 and TinyFaceDetector models), HTML5-QRCode, React-Webcam.")
    add_bullet("Document & PDF Engine: jsPDF and jspdf-autotable (vectorized client-side academic report generation).")
    add_bullet("Data Visualization: Chart.js, React-Chartjs-2, Recharts, Lucide-React icons.")
    add_bullet("Backend Core: Node.js (v22), Express 5, Mongoose 9 ODM.")
    add_bullet("Authentication & Cryptography: JSON Web Tokens (JWT), Bcrypt.js (10 salt rounds), Web Crypto API.")
    add_bullet("Database & Cloud: MongoDB Atlas multi-region replica set with in-memory fallback support.")

    add_h14("4.2 Roles and Capabilities")
    add_bullet("Admin Portal: Comprehensive institutional control. Admins manage academic Departments (e.g., FCAIT), Degree Courses / Branches (IMSCIT, BCA, MCA), Faculty Members with multi-subject assignment, Subject catalogues, Student rosters, Global Classes, and institution-wide Daily/Weekly/Monthly attendance analytics with 1-click official PDF export.")
    add_bullet("Faculty Portal: Faculty members have role-isolated access to their assigned teaching subjects, academic timetable, enrolled student rosters with Face ID status, geofenced QR session launching (custom radius and passcode), real-time attendance monitor, manual session termination, and student attendance reporting with PDF download.")
    add_bullet("Student Portal: Students securely self-register, enroll Face ID biometric embeddings, inspect daily lecture schedules, scan live classroom QR codes, submit geofenced biometrics, view attendance history, and download official attendance certificates for semester examination clearance.")

    add_h14("4.3 Attendance Verification Pipeline & Adaptive Biometrics")
    add_body(
        "Attendance verification follows an ultra-secure, multi-stage pipeline designed for seamless execution even under variable classroom lighting and low-quality device cameras:"
    )
    add_bullet("Stage 1 - Session Validation: The server validates the active status of the class session and verifies the decrypted session token against the active database document.")
    add_bullet("Stage 2 - Secret Security Passcode: The client-submitted 4-digit code must exactly match the teacher's active session passcode (HTTP 401 on mismatch).")
    add_bullet("Stage 3 - Haversine Geofence Verification: The student's device GPS coordinates (latitude, longitude, accuracy radius) are measured against the faculty device coordinates using the Haversine great-circle formula. If the calculated distance exceeds the session radius (e.g., 100m), access is denied (HTTP 403).")
    add_bullet("Stage 4 - Adaptive Multi-Pass Facial Verification: Client-side face detection first applies SSD MobileNet V1 (minConfidence: 0.20) and automatically cascades to TinyFaceDetector if lighting is dim or webcam resolution is low. The extracted 128-dimensional biometric descriptor is compared against the student's registered profile embeddings. An adaptive Euclidean distance threshold of 0.68 is enforced, returning confidence telemetry: 'Perfect Match' (< 0.45), 'High Confidence' (< 0.58), or 'Verified Match'.")
    add_bullet("Stage 5 - Idempotent Check-in & Retry Resilience: If a student re-submits in the same active session due to network latency, the server updates their coordinates and biometric telemetry and responds with HTTP 200 (alreadyMarked: true), preventing duplicate rejections.")
    add_bullet("Stage 6 - Post-Session Auto-Reconciliation: When the instructor clicks 'End Session' (POST /api/attendance/stop/:id), the server closes the QR code, marks absent all enrolled students who failed to check in, and compiles complete statistical reports.")

    add_h14("4.4 Security & Data Privacy")
    add_bullet("Zero Raw Image Storage: To protect student privacy, raw facial photographs are processed in-memory and never stored on disk or database. Only irreversible 128-dimensional mathematical floating-point embedding vectors are persisted.")
    add_bullet("Bcrypt Password Encryption: User passwords are automatically salted and hashed (10 rounds) in Mongoose pre-save middleware and strictly excluded from default query projections.")
    add_bullet("Stateless JWT Authorization: Bearer JWT tokens with 30-day expiration guard protected routes through 'protect' and 'authorize' middleware tiers.")

    add_h14("4.5 REST API Endpoints Specification")
    api_table_data = [
        ["POST", "/api/auth/register", "Public", "Register student account with initial profile data"],
        ["POST", "/api/auth/login", "Public", "Authenticate user and issue signed JWT bearer token"],
        ["GET", "/api/auth/me", "Authenticated", "Retrieve current authenticated user context"],
        ["POST", "/api/auth/face-setup", "Authenticated", "Store 128-dimensional facial embedding vector"],
        ["GET", "/api/admin/stats", "Admin", "Admin dashboard counters, attendance charts, activity stream"],
        ["GET, POST", "/api/admin/users", "Admin", "List all users (with role filters) and create user accounts"],
        ["PUT, DELETE", "/api/admin/users/:id", "Admin", "Update user role/credentials or delete user"],
        ["GET, POST", "/api/admin/departments", "Admin", "List and create academic departments (e.g. FCAIT)"],
        ["PUT, DELETE", "/api/admin/departments/:id", "Admin", "Update or delete department record"],
        ["GET, POST", "/api/admin/courses", "Admin / Auth", "List degree courses (IMSCIT, BCA) and create courses"],
        ["PUT, DELETE", "/api/admin/courses/:id", "Admin", "Update or delete degree course record"],
        ["GET, POST", "/api/classes", "Authenticated", "List classes or create new class with schedule slots"],
        ["GET, POST", "/api/subjects", "Authenticated", "List academic subjects or create new subject"],
        ["GET, POST", "/api/students", "Authenticated", "List enrolled students or register student profile"],
        ["GET", "/api/faculty-list", "Authenticated", "List faculty professors for scheduling and assignments"],
        ["POST", "/api/attendance/start", "Faculty", "Start live geofenced QR session with secret passcode"],
        ["POST", "/api/attendance/stop/:id", "Faculty / Admin", "Close active session and compile roster attendance report"],
        ["GET", "/api/attendance/session/:id", "Authenticated", "Get post-session attendance summary and distance roster"],
        ["POST", "/api/attendance/mark", "Student", "Submit QR token, passcode, GPS, and face descriptor"],
        ["GET", "/api/attendance/active-session/:classId", "Student", "Retrieve active session metadata for class"],
        ["GET", "/api/attendance", "Authenticated", "Retrieve historical attendance sessions by user role"],
        ["GET", "/api/attendance/student-stats", "Student", "Student dashboard analytics and today's schedule"],
        ["GET", "/api/attendance/student-report", "Authenticated", "Daily/Weekly/Monthly student attendance report with PDF readiness"]
    ]
    create_table(["Method", "Endpoint Route", "Access Level", "Operational Purpose"], api_table_data, [0.9, 2.3, 1.1, 2.3])

    add_h14("4.6 Project Objectives")
    add_bullet("Eliminate Proxy Attendance: Enforce multi-factor verification combining classroom proximity (GPS), physical presence (Facial recognition), and time-limited QR codes.")
    add_bullet("Automate Faculty Workflow: Replace manual paper roll calls with automated live scanning, reducing lecture attendance overhead from 15 minutes to under 60 seconds.")
    add_bullet("Hierarchical Academic Administration: Model institutional realities where 1 Department oversees Multiple Courses, Courses offer Multiple Subjects, and 1 Faculty teaches Multiple Subjects.")
    add_bullet("Audit-Ready Exporting: Empower administration, teachers, and students to generate verified Daily, Weekly, and Monthly attendance reports with 1-click PDF export.")

    add_h14("4.7 Functional and Non-Functional Requirements")
    add_h13("Table A: Functional Requirements Matrix")
    fr_data = [
        ["FR1", "Multi-Role Authentication", "Users log in securely with JWT bearer tokens and role-based redirects.", "authController.js"],
        ["FR2", "Department & Course Hierarchy", "Admins manage Departments and Degree Courses with referential links.", "adminController.js"],
        ["FR3", "Faculty Multi-Subject Assignment", "Admins assign multiple subjects to each professor in Faculty Management.", "FacultyManagement.jsx"],
        ["FR4", "Geofenced Session Hub", "Faculty starts and stops live sessions with dynamic QR and passcode.", "GenerateQR.jsx"],
        ["FR5", "Multi-Pass Biometric Verification", "Students check in using SSD MobileNet / TinyFace fallback + GPS geofence.", "ScanQR.jsx"],
        ["FR6", "Automated Post-Session Audit", "Ending a session computes present/absent roster with distance metrics.", "attendanceController.js"],
        ["FR7", "Daily/Weekly/Monthly PDF Reports", "Interactive student report generator with 1-click PDF download.", "StudentReports.jsx"]
    ]
    create_table(["ID", "Requirement", "System Description", "Implementation File"], fr_data, [0.6, 1.8, 2.5, 1.5])

    add_h13("Table B: Non-Functional Requirements Matrix")
    nfr_data = [
        ["Security", "Bcrypt 10-round password hashing, 30-day signed JWT tokens, role-restricted middleware tiers."],
        ["Privacy", "Biometric descriptors stored strictly as 128-float numeric arrays; no raw photos saved."],
        ["Reliability", "Adaptive 0.68 face distance threshold with multi-detector fallback for low-light mobile environments."],
        ["Portability", "Cross-platform responsive design supporting mobile smartphones, tablets, laptops, and desktop web."],
        ["Performance", "Sub-second database query execution via indexed MongoDB Atlas replica set."]
    ]
    create_table(["Quality Attribute", "Evidenced Implementation Details"], nfr_data, [1.8, 4.6])

    add_h14("4.8 Scope & Limitations")
    add_body(
        "Current Scope: Full institutional deployment across administrators, faculty teachers, and students covering automated attendance, timetable schedules, course-department hierarchies, and official PDF reporting."
    )
    add_body(
        "Known Limitations & Controlled Boundary Conditions:"
    )
    add_bullet("Liveness Detection: Passive facial recognition verifies 128 biometric landmarks; active 3D liveness detection (blink or nod challenge) is designated for phase 2 hardening.")
    add_bullet("Indoor GPS Signal Variance: Geofencing relies on device-reported satellite/network GPS; basement classrooms with severe signal shielding require slightly relaxed radius settings (e.g. 150m).")

    add_h14("4.9 Future Scope & Enhancements")
    add_bullet("Active Facial Liveness Challenge: Inject random eye-blink and head-turn micro-challenges during camera capture.")
    add_bullet("Offline Sync Progressive Web App (PWA): Allow students to cache encrypted check-in proofs offline when campus Wi-Fi drops.")
    add_bullet("Predictive Defaulter Analytics: Automated machine learning alerts warning students when attendance risks falling below 75%.")

    add_h14("4.10 Conclusion")
    add_body(
        "AttendAI successfully modernizes campus attendance by delivering an impenetrable yet effortless multi-factor verification pipeline. Combining dynamic rotating QR codes, secret security passcodes, Haversine geofences, and neural facial verification provides ironclad protection against proxy attendance. Furthermore, the newly integrated Department-Course-Faculty architecture and instant Daily/Weekly/Monthly PDF report export provide academic leadership with unmatched visibility into student attendance and examination eligibility."
    )

    doc.add_page_break()

    # ==================== PAGE 5: UML DIAGRAMS ====================
    add_h14("5. UML Diagrams")
    add_body(
        "The Unified Modeling Language (UML) diagrams below depict the structural architecture, relational foreign-key models, operational workflows, and state lifecycles of AttendAI."
    )

    add_h14("5.1 Use Case Diagram")
    add_body(
        "The Use Case Diagram defines all system interactions across the three system actors: Administrator, Faculty Member, and Student. Notable capabilities include 'Manage Degree Courses', 'Assign Multiple Subjects to Faculty', 'Stop Attendance Session', and 'Export Official Attendance Report (PDF)'."
    )
    add_embedded_image('page_8_img_1.jpeg', "Figure 1: Use Case Diagram - Multi-Role Actor Interactions", width=Inches(3.2), height=Inches(7.2))

    doc.add_page_break()

    add_h14("5.2 Class Diagram & Foreign Key Architecture")
    add_body(
        "The Class Diagram maps the object-oriented schemas of AttendAI, highlighting the new Course model and referential foreign keys linking Department (1 -> N) Course (1 -> N) Subject, Class (references Subject, Faculty, Course, Students), User (linked to Department, Course, and assignedSubjects), and Attendance (linking Class, Faculty, and Student records with biometric metrics)."
    )
    add_embedded_image('fig_2_class_diagram.jpeg', "Figure 2: Class Diagram - Mongoose Schema Relational Model", width=Inches(5.8))

    doc.add_page_break()

    add_h14("5.3 Swimlane Activity Diagram: Attendance Marking")
    add_body(
        "Illustrates the concurrent workflow divided into three swimlanes: Faculty, Student, and System (Express API + MongoDB Atlas). Shows QR creation, camera scanning, geofence validation, adaptive facial landmark extraction (0.68 threshold), and post-session automated absence reconciliation."
    )
    add_embedded_image('fig_3_swimlane.jpeg', "Figure 3: Swimlane Activity Diagram - End-to-End Attendance Marking Workflow", width=Inches(5.0))

    doc.add_page_break()

    add_h14("5.4 Sequence Diagram: Session Lifecycle & Verification")
    add_body(
        "Chronologically illustrates message passing between the Faculty UI, Student Mobile UI, Express Controllers, and MongoDB Atlas database during startSession, markAttendance, and stopSession."
    )
    add_embedded_image('fig_4_sequence.jpeg', "Figure 4: Sequence Diagram - Session Start, Biometric Verification, and Post-Session Close", width=Inches(5.6))

    doc.add_page_break()

    add_h14("5.5 State Chart Diagrams")
    add_body(
        "Figure 5 displays the Attendance Session lifecycle: Initial -> Active (broadcasting dynamic QR & accepting check-ins) -> Inactive (triggered via POST /api/attendance/stop/:id, executing enrolled roster reconciliation and report compilation)."
    )
    add_embedded_image('page_12_img_1.jpeg', "Figure 5: State Chart Diagram - Attendance Session Lifecycle", width=Inches(5.4))

    add_body(
        "Figure 6 displays the Student Attendance Verification state progression: Not Recorded -> Code Validated -> Geofence Validated -> Face Landmark Verified (distance <= 0.68) -> Present (biometrics saved to session records)."
    )
    add_embedded_image('page_13_img_1.jpeg', "Figure 6: State Chart Diagram - Student Biometric & Geofence Verification", width=Inches(3.2))

    doc.add_page_break()

    # ==================== PAGE 6: DATA DICTIONARY ====================
    add_h14("6. Data Dictionary")
    add_body(
        "The Data Dictionary provides complete schema specifications for all MongoDB collections in AttendAI. Each schema enforces strict types, unique indexes, validation constraints, and referential integrity using Mongoose ObjectIds."
    )

    add_h14("6.1 User Schema")
    user_schema_data = [
        ["_id", "ObjectId", "Primary Key", "Unique identifier generated by MongoDB."],
        ["name", "String", "Required, Trimmed", "Full legal name of the user."],
        ["email", "String", "Required, Unique", "Email address or unique system username (e.g., Arcenix)."],
        ["password", "String", "Required, select: false", "Bcrypt salted hash (10 rounds); omitted from default queries."],
        ["role", "String", "Enum: student, faculty, admin", "Role-based access level (defaults to student)."],
        ["idNumber", "String", "Unique, Sparse", "Enrollment number (students) or Employee ID (faculty)."],
        ["department", "ObjectId", "Ref: Department, Optional", "Academic department the user belongs to (e.g., FCAIT)."],
        ["course", "ObjectId", "Ref: Course, Optional", "Foreign key to Course schema (e.g., IMSCIT, BCA)."],
        ["assignedSubjects", "[ObjectId]", "Ref: Subject Array", "Multiple subjects taught by faculty (1 Faculty -> N Subjects)."],
        ["year", "String", "Optional", "Academic year or semester of the student."],
        ["faceEmbeddings", "[Number]", "Default: []", "128-dimensional floating-point vector extracted via face-api.js."],
        ["createdAt", "Date", "Auto (timestamps)", "Document creation timestamp."],
        ["updatedAt", "Date", "Auto (timestamps)", "Document last modification timestamp."]
    ]
    create_table(["Field Name", "Data Type", "Constraints", "Field Description"], user_schema_data, [1.3, 1.0, 1.5, 2.6])

    add_h14("6.2 Department Schema")
    dept_schema_data = [
        ["_id", "ObjectId", "Primary Key", "Unique department identifier."],
        ["name", "String", "Required, Unique", "Full department name (e.g., Faculty of Computer Applications & IT)."],
        ["code", "String", "Required, Unique, Uppercase", "Short institutional code (e.g., FCAIT, FoT, FoM)."],
        ["description", "String", "Max 500 chars", "Administrative description of the faculty/department."],
        ["createdAt", "Date", "Auto (timestamps)", "Creation timestamp."],
        ["updatedAt", "Date", "Auto (timestamps)", "Last modification timestamp."]
    ]
    create_table(["Field Name", "Data Type", "Constraints", "Field Description"], dept_schema_data, [1.3, 1.0, 1.5, 2.6])

    add_h14("6.3 Course / Branch Schema (NEW)")
    course_schema_data = [
        ["_id", "ObjectId", "Primary Key", "Unique degree course identifier."],
        ["name", "String", "Required, Trimmed", "Full degree course name (e.g., Integrated MSc (IT), BCA)."],
        ["code", "String", "Required, Unique, Uppercase", "Unique branch code (e.g., IMSCIT, BCA, MCA)."],
        ["department", "ObjectId", "Ref: Department, Required", "Foreign key linking course to parent Department (FCAIT)."],
        ["durationYears", "Number", "Default: 3, Min: 1", "Total duration of degree program in years."],
        ["totalSemesters", "Number", "Default: 6, Min: 1", "Total number of academic semesters in program."],
        ["createdAt", "Date", "Auto (timestamps)", "Document creation timestamp."],
        ["updatedAt", "Date", "Auto (timestamps)", "Document last modification timestamp."]
    ]
    create_table(["Field Name", "Data Type", "Constraints", "Field Description"], course_schema_data, [1.3, 1.0, 1.5, 2.6])

    add_h14("6.4 Subject Schema")
    subj_schema_data = [
        ["_id", "ObjectId", "Primary Key", "Unique subject identifier."],
        ["name", "String", "Required, Trimmed", "Official academic subject title (e.g., Artificial Intelligence)."],
        ["code", "String", "Required, Unique, Uppercase", "Standard course subject code (e.g., AI01, FC01)."],
        ["department", "ObjectId", "Ref: Department, Optional", "Foreign key linking subject to owning Department."],
        ["course", "ObjectId", "Ref: Course, Optional", "Foreign key linking subject to specific Degree Course."],
        ["semester", "Number", "Default: 1", "Academic semester in which subject is taught."],
        ["credits", "Number", "Required, Default: 3", "University academic credit weighting."],
        ["createdAt", "Date", "Auto (timestamps)", "Creation timestamp."],
        ["updatedAt", "Date", "Auto (timestamps)", "Last modification timestamp."]
    ]
    create_table(["Field Name", "Data Type", "Constraints", "Field Description"], subj_schema_data, [1.3, 1.0, 1.5, 2.6])

    add_h14("6.5 Class Schema")
    class_schema_data = [
        ["_id", "ObjectId", "Primary Key", "Unique class division identifier."],
        ["name", "String", "Required", "Class batch label (e.g., FY-ImscIT-AI01)."],
        ["course", "ObjectId", "Ref: Course, Optional", "Foreign key linking class division to Degree Course."],
        ["subject", "ObjectId", "Ref: Subject, Required", "Foreign key linking class division to Subject."],
        ["faculty", "ObjectId", "Ref: User, Required", "Foreign key linking class division to responsible Faculty."],
        ["students", "[ObjectId]", "Ref: User Array", "Array of enrolled student ObjectIds."],
        ["schedule", "[Object]", "Sub-document Array", "Weekly timetable schedule slots."],
        ["schedule[].dayOfWeek", "String", "Required, Enum", "Day of the week (Monday through Sunday)."],
        ["schedule[].startTime", "String", "Required, HH:MM", "Lecture start time."],
        ["schedule[].endTime", "String", "Required, HH:MM", "Lecture end time."],
        ["schedule[].room", "String", "Optional", "Classroom / laboratory room number."],
        ["schedule[].type", "String", "Default: Lecture", "Session category (Lecture, Lab, Seminar)."]
    ]
    create_table(["Field Name", "Data Type", "Constraints", "Field Description"], class_schema_data, [1.3, 1.0, 1.5, 2.6])

    add_h14("6.6 Attendance Schema & Embedded Records")
    att_schema_data = [
        ["_id", "ObjectId", "Primary Key", "Unique attendance session identifier."],
        ["class", "ObjectId", "Ref: Class, Required", "Foreign key to Class division."],
        ["faculty", "ObjectId", "Ref: User, Required", "Foreign key to Faculty professor who launched session."],
        ["date", "Date", "Required, Default: now", "Date and timestamp of class session."],
        ["sessionToken", "String", "Required, Unique", "Cryptographically random 32-byte hex QR session token."],
        ["isActive", "Boolean", "Default: true", "True while session is open; false when stopped."],
        ["attendanceCode", "String", "Required", "4-digit secret security code chosen by instructor."],
        ["location.lat", "Number", "Required", "Instructor device latitude coordinates."],
        ["location.lng", "Number", "Required", "Instructor device longitude coordinates."],
        ["radius", "Number", "Required, Default: 100", "Maximum permissible geofence distance in meters."],
        ["records", "[Object]", "Sub-document Array", "Embedded student check-in records."],
        ["records[].student", "ObjectId", "Ref: User, Required", "Foreign key to student User document."],
        ["records[].status", "String", "Enum: Present, Absent", "Final attendance status."],
        ["records[].distance", "Number", "Optional", "Measured Haversine geofence distance (meters)."],
        ["records[].faceDistance", "Number", "Optional", "Measured Euclidean distance from face-api.js."],
        ["records[].accuracy", "Number", "Optional", "Device GPS reported accuracy radius (meters)."],
        ["records[].markedAt", "Date", "Optional", "Exact timestamp when biometric verification passed."]
    ]
    create_table(["Field Name", "Data Type", "Constraints", "Field Description"], att_schema_data, [1.3, 1.0, 1.5, 2.6])

    add_h14("6.7 Relational Model & Referential Integrity")
    add_body(
        "Although MongoDB is a NoSQL document database, AttendAI rigorously implements relational referential integrity using Mongoose ObjectIds and populated relationships:"
    )
    add_bullet("Department to Courses (1 to N): Department._id is referenced by Course.department. A single department (e.g., FCAIT) encompasses multiple degree courses (IMSCIT, BCA, MCA).")
    add_bullet("Department to Faculty (1 to N): User.department references Department._id, linking professors to their parent faculty institute.")
    add_bullet("Faculty to Subjects (1 to N): User.assignedSubjects stores an array of Subject._id references, empowering 1 faculty member to teach multiple subjects across semesters.")
    add_bullet("Course to Subjects & Classes (1 to N): Subject.course and Class.course maintain degree-branch isolation, ensuring timetable slots and subject codes strictly belong to their respective curriculum.")
    add_bullet("Class to Students (M to N): Class.students stores arrays of student ObjectIds, enabling flexible multi-batch enrollments.")
    add_bullet("Attendance to Class & Faculty (N to 1): Every Attendance session references the Class._id and Faculty._id, establishing an immutable academic audit trail.")

    doc.add_page_break()

    # ==================== PAGE 7: SCREENSHOTS WITH URL ====================
    add_h14("7. All Screenshots with URL")
    add_body(
        "Base URL (Local Development): http://localhost:5173\nProduction Cloud Deployment: https://attendai-rho.vercel.app\n"
        "AttendAI provides dedicated, role-isolated portals for Administrators, Faculty Members, and Students. The table below lists all 22 system interfaces followed by the actual system screens."
    )

    add_h13("Table 6: Screen and URL Directory")
    screens_summary_data = [
        ["Admin", "Admin Dashboard", "http://localhost:5173/dashboard"],
        ["Admin", "Login Portal", "http://localhost:5173/login"],
        ["Admin", "User Management", "http://localhost:5173/dashboard/users"],
        ["Admin", "Faculty Members Management", "http://localhost:5173/dashboard/faculty"],
        ["Admin", "Courses & Degree Branches", "http://localhost:5173/dashboard/courses"],
        ["Admin", "Department Management", "http://localhost:5173/dashboard/departments"],
        ["Admin", "Subject Library", "http://localhost:5173/dashboard/subjects"],
        ["Admin", "All Institutional Classes", "http://localhost:5173/dashboard/classes"],
        ["Admin / Faculty", "Student Attendance Reports & Analytics", "http://localhost:5173/dashboard/reports"],
        ["Faculty", "Faculty Dashboard", "http://localhost:5173/dashboard"],
        ["Faculty", "My Assigned Classes", "http://localhost:5173/dashboard/my-classes"],
        ["Faculty", "Weekly Timetable Schedule", "http://localhost:5173/dashboard/timetable"],
        ["Faculty", "My Enrolled Students", "http://localhost:5173/dashboard/students"],
        ["Faculty", "Live Attendance QR Hub", "http://localhost:5173/dashboard/qr"],
        ["Faculty", "Subjects (Faculty View)", "http://localhost:5173/dashboard/subjects"],
        ["Faculty", "Session Archives (Records)", "http://localhost:5173/dashboard/records"],
        ["Student", "Student Dashboard", "http://localhost:5173/dashboard"],
        ["Student", "Self-Registration Portal", "http://localhost:5173/register"],
        ["Student", "Biometric QR Scanner", "http://localhost:5173/dashboard/scan"],
        ["Student", "My Attendance Profile & History", "http://localhost:5173/dashboard/history"],
        ["Student", "My Attendance Clearance Report (PDF)", "http://localhost:5173/dashboard/reports"],
        ["Report", "Academic Timetable Print Layout (PDF)", "http://localhost:5173/dashboard/timetable"]
    ]
    create_table(["Portal Section", "Interface Screen Name", "System Access URL"], screens_summary_data, [1.5, 2.3, 2.6])

    doc.add_page_break()

    # 7.1 Admin Side Screens
    add_h14("7.1 Admin Side Screenshots")
    add_screenshot_figure('page_23_img_1.jpeg', "7", "Admin Dashboard", "http://localhost:5173/dashboard")
    add_screenshot_figure('page_23_img_2.jpeg', "8", "Login", "http://localhost:5173/login")

    doc.add_page_break()

    add_screenshot_figure('page_24_img_1.jpeg', "9", "User Management", "http://localhost:5173/dashboard/users")
    add_screenshot_figure('page_24_img_2.jpeg', "10", "Departments", "http://localhost:5173/dashboard/departments")

    doc.add_page_break()

    add_screenshot_figure('page_25_img_1.jpeg', "11", "Subjects", "http://localhost:5173/dashboard/subjects")
    add_screenshot_figure('page_25_img_2.jpeg', "12", "All Classes", "http://localhost:5173/dashboard/classes")

    doc.add_page_break()

    # 7.2 Faculty Portal Screens
    add_h14("7.2 Faculty Portal Screenshots")
    add_screenshot_figure('page_26_img_1.jpeg', "13", "Faculty Dashboard", "http://localhost:5173/dashboard")
    add_screenshot_figure('page_26_img_2.jpeg', "14", "My Classes", "http://localhost:5173/dashboard/my-classes")

    doc.add_page_break()

    add_screenshot_figure('page_27_img_1.jpeg', "15", "Timetable", "http://localhost:5173/dashboard/timetable")
    add_screenshot_figure('page_27_img_2.jpeg', "16", "My Students", "http://localhost:5173/dashboard/students")

    doc.add_page_break()

    add_screenshot_figure('page_28_img_1.jpeg', "17", "Generate QR", "http://localhost:5173/dashboard/qr")
    add_screenshot_figure('page_28_img_2.jpeg', "18", "Subjects (Faculty)", "http://localhost:5173/dashboard/subjects")

    doc.add_page_break()

    # 7.3 Student Portal Screens
    add_h14("7.3 Student Portal Screenshots")
    add_screenshot_figure('page_29_img_1.jpeg', "19", "Student Dashboard", "http://localhost:5173/dashboard")
    add_screenshot_figure('page_29_img_2.jpeg', "20", "Register (Student)", "http://localhost:5173/register")

    doc.add_page_break()

    add_screenshot_figure('page_30_img_1.jpeg', "21", "Scan QR", "http://localhost:5173/dashboard/scan")
    add_screenshot_figure('page_30_img_2.jpeg', "22", "My History", "http://localhost:5173/dashboard/history")

    doc.add_page_break()

    # 7.4 Reports
    add_h14("7.4 Reports & Analytics Screenshots")
    add_screenshot_figure('page_31_img_1.jpeg', "23", "Attendance Records with CSV Export (Faculty)", "http://localhost:5173/dashboard/records")
    add_screenshot_figure('page_32_img_1.jpeg', "24", "Timetable Print Layout (PDF via browser print)", "http://localhost:5173/dashboard/timetable")

    doc.add_page_break()

    # ==================== PAGE 8: REFERENCES ====================
    add_h14("8. References")
    refs = [
        "[1] Express.js Documentation and Architectural Guidelines, https://expressjs.com/",
        "[2] Mongoose ODM v9 Schema and Population Reference, https://mongoosejs.com/docs/",
        "[3] MongoDB Atlas Cloud Database Documentation, https://www.mongodb.com/docs/",
        "[4] React 19 Core Documentation and Server Components, https://react.dev/",
        "[5] Vite 8 Frontend Build Tool Documentation, https://vite.dev/",
        "[6] face-api.js: JavaScript API for Face Detection and Recognition in the Browser with TensorFlow.js, https://github.com/justadudewhohacks/face-api.js/",
        "[7] JSON Web Token (JWT) Architecture, RFC 7519, Internet Engineering Task Force (IETF), https://www.rfc-editor.org/rfc/rfc7519",
        "[8] Haversine Formula for Great-Circle Distance Computations over Earth's Surface, R. W. Sinnott, Sky and Telescope, 1984.",
        "[9] jsPDF & jspdf-autotable Vectorized Client-Side Document Generation Library, https://github.com/parallax/jsPDF",
        "[10] Object Management Group (OMG), Unified Modeling Language (UML) Specification Version 2.5.1, https://www.omg.org/spec/UML/"
    ]
    for ref in refs:
        add_body(ref, space_after=6)

    out_path_1 = r"C:\Users\Abhi\Downloads\Documents\AttendAi_Complete_Project_Documentation.docx"
    out_path_2 = r"c:\xampp\htdocs\Attendai\AttendAi_Complete_Project_Documentation.docx"
    
    doc.save(out_path_1)
    doc.save(out_path_2)
    print(f"Updated document with embedded diagrams and screenshots saved successfully to:\n1. {out_path_1}\n2. {out_path_2}")

if __name__ == "__main__":
    build_complete_document()
